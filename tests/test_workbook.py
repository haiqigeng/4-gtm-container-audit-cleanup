"""The analyst workbook and its saved-file verification."""

import json

from conftest import container, finding, operation, p

from gtmaudit import source
from gtmaudit.records import Store


def _deliver(model, tmp_path, records, name="w"):
    """Records -> finalize -> workbook; returns (final, workbook path, tab rows by sheet index)."""
    from openpyxl import load_workbook

    from gtmaudit import workbook
    from gtmaudit.final import finalize
    store = Store(tmp_path / name)
    store.append("D3", records, set(model.objects), set())
    store.append("INT", [{"type": "summary", "target": "t", "naming": "n", "limits": "l"}], set(model.objects), set())
    final = finalize(model, store, [], {})
    assert not final["errors"], final["errors"]
    book = workbook.build(model, final, {}, tmp_path / f"{name}.xlsx")
    wb = load_workbook(book)
    return final, book, [list(ws.iter_rows(min_row=2, values_only=True)) for ws in wb.worksheets]


def _work(i, op_ids=None, **extra):
    return {**finding(id=f"D3-F0{i}", area="A22", operations=op_ids or [f"D3-OP0{i}"], **extra),
            "class": "optimisation", "basis": ""}


def _html(value, obj="tag:3"):
    return {"do": "set", "object": obj, "path": "parameter:html.value", "value": value}


def test_workbook_places_each_change_once_and_keeps_text_literal(model, tmp_path):
    from openpyxl import load_workbook

    from gtmaudit import workbook
    from gtmaudit.final import finalize
    code = "<script>" + "x" * (workbook.MAX_CELL + 10) + "</script>"
    store = Store(tmp_path)
    store.append("D3", [
        {**finding(id="D3-F01", area="A22", operations=["D3-OP01"], title="=SUM(A1:A2)"), "class": "optimisation",
         "basis": ""},
        {**finding(id="D3-F02", area="A22", severity="Low", operations=["D3-OP01"]), "class": "optimisation", "basis": ""},
        operation(id="D3-OP01", findings=["D3-F01", "D3-F02"], actions=[
            {"do": "set", "object": "tag:3", "path": "parameter:html.value", "value": code}])],
        set(model.objects), set())
    store.append("INT", [{"type": "summary", "target": "t", "naming": "n", "limits": "l"}], set(model.objects), set())
    final = finalize(model, store, [], {})
    book = workbook.build(model, final, {}, tmp_path / "w.xlsx")
    wb = load_workbook(book)
    rows = list(wb["2 Findings and actions"].iter_rows(min_row=2))
    assert rows[0][4].data_type == "s" and rows[0][4].value == "=SUM(A1:A2)"
    assert "1. " in rows[0][8].value and 'Done by the change "Fire per event" in #1' in rows[1][8].value
    assert 'the new code in tab 4 Custom code › Tag "Meta - base"' in rows[0][8].value
    parts = [r for r in wb["4 Custom code"].iter_rows(min_row=2, values_only=True) if r[0].startswith('Tag "Meta - base"')]
    assert len(parts) == 2 and "".join(r[5] for r in parts) == code
    assert workbook.verify(book, model, final, {}) == []


def test_code_tab_shows_like_for_like_material_in_complete_parts(tmp_path):
    from openpyxl import load_workbook

    from gtmaudit import source, workbook
    from gtmaudit.final import finalize
    raw = container()
    long_code = "<script>" + "y" * (workbook.MAX_CELL + 5) + "</script>"
    raw["containerVersion"]["tag"][2]["parameter"] = [p("html", long_code)]
    export = tmp_path / "export.json"
    export.write_text(json.dumps(raw), encoding="utf-8")
    model = source.load(export)
    new_file = raw["containerVersion"]["customTemplate"][0]["templateData"].replace("vendor.example", "cdn.example")
    store = Store(tmp_path)
    store.append("D3", [{**finding(id="D3-F01", area="A22", objects=["template:40"], operations=["D3-OP01"]),
                         "class": "optimisation", "basis": ""},
                        operation(id="D3-OP01", findings=["D3-F01"], actions=[
                            {"do": "set", "object": "template:40", "path": "templateData", "value": new_file}])],
                 set(model.objects), set())
    store.append("INT", [{"type": "summary", "target": "t", "naming": "n", "limits": "l"}], set(model.objects), set())
    final = finalize(model, store, [], {})
    book = workbook.build(model, final, {}, tmp_path / "w.xlsx")
    wb = load_workbook(book)
    code = list(wb["4 Custom code"].iter_rows(min_row=2, values_only=True))
    template = [r for r in code if r[0].startswith('Template "Vendor template"')]
    assert template[0][4] == raw["containerVersion"]["customTemplate"][0]["templateData"] and template[0][5] == new_file
    assert "imported with More actions › Import" in next(wb["2 Findings and actions"].iter_rows(min_row=2, values_only=True))[8]
    meta = [r for r in code if r[0].startswith('Tag "Meta - base"')]
    assert len(meta) == 2 and "".join(r[4] for r in meta) == long_code
    assert workbook.verify(book, model, final, {}) == []
    wb["4 Custom code"]["F2"] = "// replaced"
    wb.save(book)
    assert any("differs from final.json" in q for q in workbook.verify(book, model, final, {}))


def test_each_code_change_shows_its_own_material(model, tmp_path):
    from openpyxl import load_workbook

    from gtmaudit import workbook
    from gtmaudit.final import finalize
    source_code = model.get("tag:3").raw["parameter"][0]["value"]
    cleanup, repair = "<script>fbq('init','123456');</script>", "<script>fbq('init','123456');fbq('track','X');</script>"
    first = operation(id="D3-OP01", findings=["D3-F01"], title="Remove the page view call",
                      actions=[{"do": "set", "object": "tag:3", "path": "parameter:html.value", "value": cleanup}])
    second = operation(id="D3-OP02", findings=["D3-F02"], title="Send the X event", depends_on=["D3-OP01"],
                       actions=[{"do": "set", "object": "tag:3", "path": "parameter:html.value", "value": repair}])
    findings = [{**finding(id=f"D3-F0{i}", area="A22", operations=[f"D3-OP0{i}"]), "class": "optimisation", "basis": ""}
                for i in (1, 2)]

    def build(records, name):
        store = Store(tmp_path / name)
        store.append("D3", records, set(model.objects), set())
        store.append("INT", [{"type": "summary", "target": "t", "naming": "n", "limits": "l"}], set(model.objects), set())
        final = finalize(model, store, [], {})
        return final, workbook.build(model, final, {}, tmp_path / f"{name}.xlsx")
    final, book = build(findings + [first, second], "a")
    _, swapped = build(findings + [second, first], "b")
    rows = [r for r in load_workbook(book)["4 Custom code"].iter_rows(min_row=2, values_only=True)
            if r[0].startswith('Tag "Meta - base"')]
    assert [(r[0], r[4], r[5]) for r in rows] == [
        ('Tag "Meta - base" · Remove the page view call', source_code, cleanup),
        ('Tag "Meta - base" · Send the X event', cleanup, repair)]
    assert rows == [r for r in load_workbook(swapped)["4 Custom code"].iter_rows(min_row=2, values_only=True)
                    if r[0].startswith('Tag "Meta - base"')]
    tab2 = {r[13].split("\n")[0]: r[8] for r in load_workbook(book)["2 Findings and actions"].iter_rows(min_row=2, values_only=True)}
    assert "› Tag \"Meta - base\" · Send the X event" in tab2["D3-F02"]
    assert workbook.verify(book, model, final, {}) == []
    wb = load_workbook(book)
    cells = [c for c in wb["4 Custom code"]["F"] if c.value in (cleanup, repair)]
    cells[0].value, cells[1].value = cells[1].value, cells[0].value
    wb.save(book)
    assert any("is not that action's own before/after code" in q for q in workbook.verify(book, model, final, {}))


def test_two_code_writes_in_one_operation_keep_their_own_material(model, tmp_path):
    from gtmaudit import workbook
    source_code = model.get("tag:3").raw["parameter"][0]["value"]
    first, second = "<script>fbq('init','1');</script>", "<script>fbq('init','2');</script>"
    final, book, tabs = _deliver(model, tmp_path, [_work(1), operation(id="D3-OP01", findings=["D3-F01"],
                                                               actions=[_html(first), _html(second)])])
    code = {r[7]: (r[4], r[5]) for r in tabs[3] if r[7]}
    assert code == {"D3-OP01 · 1": (source_code, first), "D3-OP01 · 2": (first, second)}
    todo = tabs[1][0][8]
    assert "(Ref D3-OP01 · 1)" in todo and "(Ref D3-OP01 · 2)" in todo
    assert workbook.verify(book, model, final, {}) == []


def test_steps_use_the_state_their_prerequisites_leave(model, tmp_path):
    rename = operation(id="D3-OP01", findings=["D3-F01"],
                       actions=[{"do": "rename", "object": "tag:3", "name": "Meta - Base pixel"}])
    edit = operation(id="D3-OP02", findings=["D3-F02"], depends_on=["D3-OP01"],
                     actions=[{"do": "set", "object": "tag:3", "path": "notes", "value": "kept"}])
    create = operation(id="D3-OP03", findings=["D3-F03"], actions=[{"do": "create", "layer": "tag", "ref": "new:tag:echo",
        "object": {"name": "Echo constant", "type": "html", "firingTriggerId": ["11"],
                   "parameter": [p("html", "<script>console.log({{Unused constant}})</script>")]}}])
    remap = operation(id="D3-OP04", findings=["D3-F04"], depends_on=["D3-OP03"],
                      actions=[{"do": "remap", "from": "variable:22", "to": "variable:20"}])
    _, _, tabs = _deliver(model, tmp_path, [_work(i) for i in (1, 2, 3, 4)] + [rename, edit, create, remap])
    todo = {r[13].split("\n")[0]: r[8] for r in tabs[1]}
    assert 'Tag "Meta - Base pixel": set Notes to "kept"' in todo["D3-F02"]
    assert 'Tag "Echo constant"' in todo["D3-F04"].split("everywhere it is used:")[1]


def test_a_row_lists_prerequisites_first(model, tmp_path):
    from gtmaudit import workbook
    base = operation(id="D3-OP01", findings=["D3-F01"], title="Prepare", actions=[_html("<script>a()</script>")])
    then = operation(id="D3-OP02", findings=["D3-F01"], title="Finish", depends_on=["D3-OP01"],
                     actions=[_html("<script>b()</script>")])
    final, book, tabs = _deliver(model, tmp_path, [_work(1, ["D3-OP02", "D3-OP01"]), then, base])
    todo = tabs[1][0][8]
    assert todo.index("Change 1: Prepare") < todo.index("Change 2: Finish (after change 1)")
    assert tabs[1][0][13].split("\n")[1:] == ["D3-OP01", "D3-OP02"]
    assert workbook.verify(book, model, final, {}) == []


def test_identical_titles_and_renamed_variables_keep_exact_material(tmp_path):
    from openpyxl import load_workbook

    from gtmaudit import workbook
    raw = container()
    raw["containerVersion"]["tag"][2]["parameter"] = [p("html", "<script>fbq('track','P',{v:{{DLV - value}}})</script>")]
    export = tmp_path / "export.json"
    export.write_text(json.dumps(raw), encoding="utf-8")
    model = source.load(export)
    rename = operation(id="D3-OP01", findings=["D3-F01"], title="Same title",
                       actions=[{"do": "rename", "object": "variable:20", "name": "DLV - order value"}])
    edit = operation(id="D3-OP02", findings=["D3-F02"], title="Same title", depends_on=["D3-OP01"],
                     actions=[_html("<script>fbq('track','Q',{v:{{DLV - order value}}})</script>")])
    again = operation(id="D3-OP03", findings=["D3-F03"], title="Same title", depends_on=["D3-OP02"],
                      actions=[_html("<script>fbq('track','R')</script>")])
    final, book, tabs = _deliver(model, tmp_path, [_work(i) for i in (1, 2, 3)] + [rename, edit, again])
    code = {r[7]: (r[4], r[5]) for r in tabs[3] if r[7]}
    assert code["D3-OP02 · 1"][0] == "<script>fbq('track','P',{v:{{DLV - order value}}})</script>"
    assert code["D3-OP03 · 1"][0] == code["D3-OP02 · 1"][1]
    assert workbook.verify(book, model, final, {}) == []
    wb = load_workbook(book)
    rows = [row for row in wb["4 Custom code"].iter_rows(min_row=2) if row[7].value]
    rows[0][7].value, rows[1][7].value = rows[1][7].value, rows[0][7].value
    wb.save(book)
    assert any("before/after code" in q for q in workbook.verify(book, model, final, {}))
