"""The command line end to end, in both workbook languages."""

import json
import subprocess
import sys

import pytest
from conftest import SCRIPTS, finding, operation

from gtmaudit import rules


def run(*args, stdin=None):
    return subprocess.run([sys.executable, "-X", "utf8", "-B", str(SCRIPTS / "gtm.py"), *map(str, args)],
                          input=stdin, capture_output=True, text=True, encoding="utf-8")


@pytest.mark.parametrize("language", ["en", "fr"])
def test_end_to_end(export, tmp_path, language):
    ctx = tmp_path / "ctx.json"
    ctx.write_text(json.dumps({"source_complete": True, "language": language, "answers": {}}), encoding="utf-8")
    out = tmp_path / "out"
    assert run("prepare", export, "--context", ctx, "--out", out).returncode == 0
    candidates = json.loads((out / "candidates.json").read_text(encoding="utf-8"))["candidates"]
    cap = next(c["id"] for c in candidates if c["rule"] == "F04")
    assert run("add", out, "--worker", "D4", "--file", "-",
               stdin=json.dumps([finding(candidates=[cap]), operation()])).returncode == 0
    shared = operation(id="D2-OP01", findings=["D2-F01"], actions=[
        {"do": "create", "layer": "variable", "ref": "new:variable:purchase-settings", "object": {
            "name": "GTES - purchase", "type": "gtes", "parameter": [{"type": "LIST", "key": "eventSettingsTable",
                                                                      "list": []}]}},
        {"do": "set", "object": "tag:2", "path": "parameter:eventSettingsVariable.value", "value": "{{GTES - purchase}}"}])
    shared_finding = {**finding(id="D2-F01", area="A15", operations=["D2-OP01"]), "class": "optimisation", "basis": ""}
    r = run("add", out, "--worker", "D2", "--file", "-", stdin=json.dumps([shared_finding, shared]))
    assert r.returncode == 0, r.stderr
    assert run("check", out, "--worker", "D2").returncode in (0, 1)
    for domain in rules.DOMAINS:
        ids = [c["id"] for c in candidates if c["domain"] == domain and c["id"] != cap]
        if ids:
            r = run("add", out, "--worker", domain, "--file", "-",
                    stdin=json.dumps({"type": "dismiss", "candidates": ids, "reason": "test"}))
            assert r.returncode == 0, r.stderr
    assert run("finalize", out).returncode == 1
    # Integrator: replace D4's finding and operation with new INT ids, then summarise.
    int_finding = finding(id="INT-F01", operations=["INT-OP01"], candidates=[cap])
    int_operation = operation(id="INT-OP01", findings=["INT-F01"], depends_on=["D2-OP01"])
    r = run("add", out, "--worker", "INT", "--file", "-", stdin=json.dumps([
        {"type": "decision", "action": "replace", "ids": ["D4-F01"], "reason": "r", "record": int_finding},
        {"type": "decision", "action": "replace", "ids": ["D4-OP01"], "reason": "r", "record": int_operation},
        {"type": "summary", "target": "t", "naming": "n", "limits": "l"}]))
    assert r.returncode == 0, r.stderr
    assert run("finalize", out).returncode == 0
    # Verifier reads final.json ids, including ids that exist only through integrator decisions.
    r = run("add", out, "--worker", "VER", "--file", "-", stdin=json.dumps([
        {"type": "verdict", "target": "INT-OP01", "verdict": "fix", "reason": "r", "change": "c"},
        {"type": "verdict", "target": "INT-F01", "verdict": "confirm", "reason": "r"}]))
    assert r.returncode == 0, r.stderr
    r = run("check", out, "--worker", "VER")
    assert r.returncode == 0, r.stdout
    r = run("finalize", out)
    assert r.returncode == 1 and "not yet handled" in r.stdout
    run("add", out, "--worker", "INT", "--file", "-",
        stdin=json.dumps({"type": "decision", "action": "resolve", "ids": ["INT-OP01"], "reason": "kept, see note"}))
    r = run("finalize", out)
    assert r.returncode == 0, r.stdout
    r = run("build", out)
    assert r.returncode == 0, r.stdout + r.stderr
    from openpyxl import load_workbook
    book = out / "gtm-container-audit.xlsx"
    wb = load_workbook(book)
    fr = language == "fr"
    assert wb.sheetnames == (["1 Synthèse", "2 Constats et actions", "3 Questions", "4 Code personnalisé", "5 Inventaire"]
                             if fr else ["1 Summary", "2 Findings and actions", "3 Questions", "4 Custom code",
                                         "5 Inventory"])
    rows = list(wb[wb.sheetnames[1]].iter_rows(min_row=2, values_only=True))
    cap_row = next(r for r in rows if str(r[13]).startswith("INT-F01"))
    todo = cap_row[8]
    assert ("Paramètres avancés › Options de déclenchement de la balise" if fr else
            "Advanced settings › Tag firing options") in todo
    assert ("Une fois par événement" if fr else "Once per event") in todo
    assert "tag:4" not in todo and "ONCE_PER_EVENT" not in todo
    assert ("Attend la réponse à Q1" if fr else "Needs answer to Q1") in cap_row[9]
    questions = list(wb[wb.sheetnames[2]].iter_rows(min_row=2, values_only=True))
    assert questions[0][:2] == ("Q1", "Does the vendor expect one call per add-to-cart?")
    inventory = [r[0] for r in wb[wb.sheetnames[4]].iter_rows(min_row=2, values_only=True)]
    assert len(inventory) == len(set(inventory))
    from gtmaudit import source as src
    from gtmaudit import workbook
    final = json.loads((out / "final.json").read_text(encoding="utf-8"))
    model = src.load(export)
    assert workbook.verify(book, model, final, {"language": language}) == []
    ws = wb[wb.sheetnames[1]]
    ws.cell(row=cap_row[0] + 1, column=9).value = "changed"
    wb.save(book)
    assert any("differs from final.json" in p for p in workbook.verify(book, model, final, {"language": language}))


def test_context_validation(export, tmp_path):
    ctx = tmp_path / "ctx.json"
    ctx.write_text(json.dumps({"source_complete": True, "do_not_touch": ["tag:404"]}), encoding="utf-8")
    r = run("prepare", export, "--context", ctx, "--out", tmp_path / "out")
    assert r.returncode == 2 and "unknown objects" in r.stderr
