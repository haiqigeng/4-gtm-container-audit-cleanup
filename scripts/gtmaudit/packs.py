"""Readable evidence packs: one whole-container summary plus one pack per review domain."""

from __future__ import annotations

import csv
import io
import json
from collections import defaultdict
from pathlib import Path

from . import VERSION
from .privacy import redact
from .render import render, variable_source
from .rules import DOMAINS, Candidate, Result
from .source import (
    CONSENT_VAR_RE,
    GOOGLE_TAG_TYPES,
    Model,
    Obj,
    code_of,
    conditions,
    consent_controls,
    has_consent_evidence,
    is_cmp_tag,
    pval,
    reads_closure,
    tag_events,
    vendors_of,
)

PAGE_LIMIT = 100_000
GOOGLE_VARIABLE_TYPES = {"gtes", "gtcs", "awec", "gas"}


def _one_line(model: Model, o: Obj) -> str:
    if o.layer == "tag":
        events = sorted(tag_events(model, o))
        ctl = consent_controls(model, o)
        consent = "+".join(n for n in ("positive_trigger", "blocker", "additional_check") if ctl[n]) or "-"
        return (f"- {o.key} {o.type} \"{o.name}\" {'PAUSED ' if o.paused else ''}"
                f"events={','.join(events[:3])}{'...' if len(events) > 3 else ''} consent={consent} "
                f"vendor={'/'.join(vendors_of(model, o))}")
    if o.layer == "trigger":
        return f"- {o.key} {o.type} \"{o.name}\" used_by={len(model.consumers(o.key))}"
    if o.layer == "variable":
        path = pval(o.raw, "name") if o.type == "v" else ""
        return f"- {o.key} {o.type} \"{o.name}\"{' path=' + path if path else ''} used_by={len(model.consumers(o.key))}"
    return f"- {o.key} {o.type} \"{o.name}\""


def render_plan(path: Path) -> str:
    """Approved tracking plan as readable, masked text (CSV, XLSX, JSON or plain text)."""
    suffix = path.suffix.lower()
    blocks = [f"# Tracking plan ({path.name})", "Supplied by the owner. Compare configured events with it.", ""]
    if suffix in (".csv", ".tsv"):
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        blocks += [" | ".join(row) for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)
                    if any(v not in (None, "") for v in row)]
            if rows:
                blocks += [f"## {ws.title}"] + rows + [""]
    elif suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("events") or data.get("requirements") or [data]
        for r in rows:
            blocks.append(" | ".join(f"{k}={v}" for k, v in r.items()) if isinstance(r, dict) else json.dumps(r))
    else:
        blocks.append(path.read_text(encoding="utf-8"))
    return redact("\n".join(blocks)) + "\n"


def summary(model: Model, result: Result, context: dict, has_plan: bool = False) -> str:
    container = model.cv.get("container") or {}
    counts = defaultdict(int)
    for o in model.objects.values():
        counts[o.layer] += 1
    lines = [
        f"# Container summary ({container.get('publicId', '?')} \"{container.get('name', '')}\")",
        f"skill v{VERSION} · source sha256 {model.sha256[:16]}… · version {model.cv.get('containerVersionId', '?')}",
        "counts: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())) + f", built-ins {len(model.builtins)}",
        "candidates per domain: " + ", ".join(
            f"{d} {sum(1 for c in result.candidates if c.domain == d)}" for d in DOMAINS),
        "", "## Supplied context (anything absent is unknown)",
        f"- protected objects (do_not_touch): {context.get('do_not_touch') or 'none'}",
        f"- Advanced Consent Mode approved destinations: {context.get('consent_mode_destinations') or 'none'}",
        f"- server hosts approved to own consent gating: {context.get('consent_owned_hosts') or 'none'}",
        f"- naming convention: {context.get('naming_convention') or 'not supplied'}",
        f"- tracking plan: {'packs/plan.md' if has_plan else 'not supplied'}",
        "- answers: " + (json.dumps(context["answers"], ensure_ascii=False) if context.get("answers") else "none"),
        "", "## Engine facts",
    ]
    for name, value in result.facts.items():
        lines.append(f"- {name}: {json.dumps(value, ensure_ascii=False)}")
    for layer in ("tag", "trigger", "variable", "template", "folder", "zone"):
        objs = sorted(model.layer(layer), key=lambda o: int(o.id) if o.id.isdigit() else 0)
        if objs:
            lines += ["", f"## {layer}s ({len(objs)})"] + [_one_line(model, o) for o in objs]
    return "\n".join(lines) + "\n"


def _candidate_block(model: Model, c: Candidate) -> str:
    lines = [f"### {c.id} · {c.rule} · {c.severity} · {c.title}"]
    for item in c.items:
        detail = f" — {item.detail}" if item.detail else ""
        lines.append(f"- {', '.join(model.label(k) for k in item.objects[:6])}"
                     f"{' …+' + str(len(item.objects) - 6) if len(item.objects) > 6 else ''}: {item.title}{detail}")
    return "\n".join(lines)


def domain_objects(model: Model, domain: str, candidates: list[Candidate]) -> list[Obj]:
    keys: list[str] = []
    for c in candidates:
        keys += c.objects
    tags = model.layer("tag")
    if domain == "D1":
        for t in tags:
            ctl = consent_controls(model, t)
            if is_cmp_tag(model, t) or ctl["blocker"] or ctl["positive_trigger"] or t.type == "awud":
                keys.append(t.key)
        keys += [tr.key for tr in model.layer("trigger") if has_consent_evidence(model, tr.id)]
        keys += [v.key for v in model.layer("variable")
                 if CONSENT_VAR_RE.search(v.name) or v.type == "awec" or CONSENT_VAR_RE.search(pval(v.raw, "name"))]
    elif domain == "D2":
        google = [t for t in tags if t.type in GOOGLE_TAG_TYPES]
        keys += [t.key for t in google]
        for t in google:
            keys += [k for k, kind in model.dependencies(t.key) if k.startswith("variable:")]
        keys += [v.key for v in model.layer("variable") if v.type in GOOGLE_VARIABLE_TYPES]
    elif domain == "D3":
        keys += [t.key for t in tags if t.type not in GOOGLE_TAG_TYPES]
        keys += [v.key for v in model.layer("variable") if v.type == "jsm"]
        keys += [t.key for t in model.layer("template")]
    elif domain == "D4":
        keys += [o.key for o in model.layer("trigger") + model.layer("variable") + model.layer("zone")
                 if not code_of(o)]
    seen, out = set(), []
    for k in keys:
        if k in model.objects and k not in seen:
            seen.add(k)
            out.append(model.objects[k])
    order = {"tag": 0, "trigger": 1, "variable": 2, "template": 3, "zone": 4, "folder": 5, "gtagConfig": 6}
    return sorted(out, key=lambda o: (order.get(o.layer, 9), int(o.id) if o.id.isdigit() else 0))


def family_view(model: Model) -> list[str]:
    """D5 blind pass: the whole container by family, without any engine candidate."""
    blocks: list[str] = []
    families = defaultdict(list)
    for t in model.layer("tag"):
        vendor = "/".join(vendors_of(model, t))
        events = sorted(tag_events(model, t))
        families[vendor].append((t, events))
    for vendor, members in sorted(families.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        lines = [f"## {vendor} ({len(members)} tags)"]
        for t, events in sorted(members, key=lambda m: (m[0].type, m[1], m[0].name)):
            key_params = {k: pval(t.raw, k) for k in ("eventName", "measurementIdOverride", "measurementId", "tagId",
                                                       "conversionId", "conversionLabel", "eventType", "pixelId",
                                                       "standardEventName") if pval(t.raw, k)}
            html = pval(t.raw, "html")
            sig = f" code={len(html)}ch" if html else ""
            lines.append(f"- {t.key} {t.type} \"{t.name}\"{' PAUSED' if t.paused else ''} events={','.join(events[:4])} "
                         f"option={t.raw.get('tagFiringOption', '-')} params={json.dumps(key_params, ensure_ascii=False)}"
                         f" blockers={t.raw.get('blockingTriggerId', [])}{sig}")
        blocks.append("\n".join(lines))
    by_kind = defaultdict(list)
    for tr in model.layer("trigger"):
        shape = "; ".join(f"{c['arg0']} {c['op']} {c['arg1']!r}" for c in conditions(tr.raw))[:160]
        by_kind[tr.type].append(f"- {tr.key} \"{tr.name}\" {shape} used_by={len(model.consumers(tr.key))}")
    for kind, rows in sorted(by_kind.items()):
        blocks.append("\n".join([f"## Triggers {kind} ({len(rows)})"] + sorted(rows)))
    by_var = defaultdict(list)
    for v in model.layer("variable"):
        detail = pval(v.raw, "name") or pval(v.raw, "value") or (f"code {len(code_of(v))}ch" if code_of(v) else "")
        by_var[v.type].append(f"- {v.key} \"{v.name}\" {detail[:120]} used_by={len(model.consumers(v.key))}")
    for kind, rows in sorted(by_var.items()):
        blocks.append("\n".join([f"## Variables {kind} ({len(rows)})"] + sorted(rows)))
    return blocks


def _paginate(text_blocks: list[str], header: str) -> list[str]:
    pages, current = [], header
    for block in text_blocks:
        if len(current.encode("utf-8")) + len(block.encode("utf-8")) > PAGE_LIMIT and current != header:
            pages.append(current)
            current = header + "(continued)\n\n"
        current += block + "\n\n"
    pages.append(current)
    return pages


def write_packs(model: Model, result: Result, context: dict, out: Path, plan: str | None = None) -> dict[str, list[str]]:
    packs = out / "packs"
    packs.mkdir(parents=True, exist_ok=True)
    written: dict[str, list[str]] = {}
    (packs / "summary.md").write_text(summary(model, result, context, plan is not None), encoding="utf-8")
    written["summary"] = ["summary.md"]
    if plan is not None:
        (packs / "plan.md").write_text(plan, encoding="utf-8")
        written["plan"] = ["plan.md"]
    for domain, title in DOMAINS.items():
        candidates = [c for c in result.candidates if c.domain == domain]
        header = (f"# {domain} {title}\nDispose every candidate below: confirm it in a finding's `candidates` "
                  f"or `dismiss` it with a reason. Look beyond the candidates within your areas.\n\n")
        blocks = ["## Candidates"] + [_candidate_block(model, c) for c in candidates] + ["## Objects"]
        objects = domain_objects(model, domain, candidates)
        blocks += [render(model, o, chains=False) for o in objects]
        sources = reads_closure(model, [o.key for o in objects])
        if sources:
            lines = [f"- {variable_source(model, k)}" for k in sources]
            blocks.append("\n".join(["## Value sources (each variable once: source and the variables it reads)"] + lines))
        pages = _paginate(blocks, header)
        names = []
        for i, page in enumerate(pages, 1):
            name = f"{domain}.md" if len(pages) == 1 else f"{domain}-{i}.md"
            (packs / name).write_text(page, encoding="utf-8")
            names.append(name)
        written[domain] = names
    blind = _paginate(family_view(model), "# D5 family view (candidate-blind first pass)\n"
                      "Read this first and record discoveries with `add` before opening D5.md.\n\n")
    names = []
    for i, page in enumerate(blind, 1):
        name = "D5-blind.md" if len(blind) == 1 else f"D5-blind-{i}.md"
        (packs / name).write_text(page, encoding="utf-8")
        names.append(name)
    written["D5-blind"] = names
    return written
