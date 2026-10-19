"""Analyst workbook built only from final.json: five tabs with one job each, then reopened and verified.

1 Summary (where do I stand), 2 Findings and actions (what is wrong and what to do), 3 Questions (what to ask),
4 Custom code (what the scripts are and what replaces them), 5 Inventory (what happens to each object).
Each fact has one home tab; other tabs point to it by row number. GTM wording comes from steps.Interface.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from . import VERSION
from .privacy import EMAIL_RE, redact
from .source import TEMPLATE_CODE, Model, code_of, strip_template, template_sections
from .operations import execution_order, trace
from .steps import Interface

LABELS = {
    "en": {
        "sheets": ("1 Summary", "2 Findings and actions", "3 Questions", "4 Custom code", "5 Inventory"),
        "findings": ("#", "Priority", "Type", "Topic", "Problem", "GTM objects", "What is wrong", "Why it matters",
                     "What to do", "Status", "Check and undo", "Decision", "Comment", "Ref"),
        "questions": ("Q#", "Question", "Problems", "Answer", "Answered by and date"),
        "code": ("Object", "Type", "Status", "Findings", "Current code", "New code", "Changed in", "Ref"),
        "inventory": ("Object", "GTM type", "Folder", "Status", "Outcome", "New name", "Findings"),
        "summary": ("Container", "Export version", "Audit date", "Scope and limits", "Target", "Naming",
                    "Open questions", "How to use this file"),
        "grid": "Priority", "total": "Total",
        "priority": {"Critical": "Critical", "High": "High", "Medium": "Medium", "Low": "Low", "None": "-"},
        "type": {"defect": "Fix", "optimisation": "Improve", "decision": "Owner decides",
                 "evidence_limit": "Check first", "keep": "Keep as is"},
        "status": {"ready": "Ready to apply", "conditional": "Needs answer to {q}", "blocked": "Blocked until {q} is answered",
                   "decision": "Owner decides: {q}", "evidence_limit": "Check first: {q}", "keep": "No change",
                   "after": "after {rows}"},
        "object_status": {"active": "Active", "paused": "Paused", "proposed": "Proposed"},
        "outcome": {"remove": "Remove", "rename": "Rename", "rename_change": "Rename and change", "change": "Change",
                    "create": "Create", "finding": "Keep (see finding)", "none": "Keep"},
        "change": "Change {n}: {title}", "after_change": "after change {n}", "before": "Before: {text}", "done_by": "Done by the change \"{title}\" in {row}",
        "check": "Check: {text}", "undo": "Undo: {text}", "part": "part {n}/{total}",
        "new_code": "the new code in tab {ref}",
        "import": "the template file in tab {ref}, imported with More actions › Import",
        "decisions": "Apply,Reject,Later",
        "boundary": ("Static audit of one web-container export. No GTM change was made and no exported code was "
                     "executed. Runtime behaviour, server containers, what vendors receive and legal compliance are "
                     "outside this audit; the proposed target is validated statically, not tested on the site."),
        "howto": ("1. Read tab 2 from the top: rows are sorted by priority.\n"
                  "2. For each row, read What is wrong and Why it matters, then choose Apply, Reject or Later in "
                  "Decision.\n"
                  "3. Send tab 3 to the site owner. Rows marked Needs answer wait for those answers.\n"
                  "4. Apply approved changes in a new GTM workspace, following What to do in row order, then run the "
                  "checks in Preview before publishing. Tab 4 holds the code to paste; tab 5 says what happens to "
                  "every object."),
    },
    "fr": {
        "sheets": ("1 Synthèse", "2 Constats et actions", "3 Questions", "4 Code personnalisé", "5 Inventaire"),
        "findings": ("#", "Priorité", "Type", "Thème", "Problème", "Objets GTM", "Ce qui ne va pas",
                     "Pourquoi c'est important", "Que faire", "Statut", "Vérifier et annuler", "Décision", "Commentaire",
                     "Réf."),
        "questions": ("Q#", "Question", "Problèmes", "Réponse", "Répondu par et date"),
        "code": ("Objet", "Type", "Statut", "Constats", "Code actuel", "Nouveau code", "Modifié en", "Réf."),
        "inventory": ("Objet", "Type GTM", "Dossier", "Statut", "Résultat", "Nouveau nom", "Constats"),
        "summary": ("Conteneur", "Version exportée", "Date de l'audit", "Périmètre et limites", "Cible", "Nommage",
                    "Questions ouvertes", "Mode d'emploi"),
        "grid": "Priorité", "total": "Total",
        "priority": {"Critical": "Critique", "High": "Haute", "Medium": "Moyenne", "Low": "Basse", "None": "-"},
        "type": {"defect": "Corriger", "optimisation": "Améliorer", "decision": "Décision du client",
                 "evidence_limit": "À vérifier", "keep": "Conserver"},
        "status": {"ready": "Prêt à appliquer", "conditional": "Attend la réponse à {q}",
                   "blocked": "Bloqué jusqu'à la réponse à {q}", "decision": "Décision du client : {q}",
                   "evidence_limit": "À vérifier d'abord : {q}", "keep": "Aucun changement", "after": "après {rows}"},
        "object_status": {"active": "Actif", "paused": "En pause", "proposed": "Proposé"},
        "outcome": {"remove": "Supprimer", "rename": "Renommer", "rename_change": "Renommer et modifier",
                    "change": "Modifier", "create": "Créer", "finding": "Conserver (voir constat)", "none": "Conserver"},
        "change": "Modification {n} : {title}", "after_change": "après la modification {n}",
        "before": "Avant : {text}",
        "done_by": "Fait par la modification « {title} » en {row}",
        "check": "Vérifier : {text}", "undo": "Annuler : {text}", "part": "partie {n}/{total}",
        "new_code": "le nouveau code de l'onglet {ref}",
        "import": "le fichier de modèle de l'onglet {ref}, importé via Plus d'actions › Importer",
        "decisions": "Appliquer,Rejeter,Plus tard",
        "boundary": ("Audit statique d'un export de conteneur web. Aucune modification GTM n'a été faite et aucun code "
                     "exporté n'a été exécuté. Le comportement réel, les conteneurs serveur, ce que reçoivent les "
                     "outils et la conformité juridique sont hors périmètre ; la cible proposée est validée "
                     "statiquement, pas testée sur le site."),
        "howto": ("1. Lisez l'onglet 2 depuis le haut : les lignes sont triées par priorité.\n"
                  "2. Pour chaque ligne, lisez Ce qui ne va pas et Pourquoi c'est important, puis choisissez "
                  "Appliquer, Rejeter ou Plus tard dans Décision.\n"
                  "3. Envoyez l'onglet 3 au responsable du site. Les lignes « Attend la réponse » dépendent de ces "
                  "réponses.\n"
                  "4. Appliquez les modifications validées dans un nouvel espace de travail GTM, en suivant Que faire "
                  "dans l'ordre des lignes, puis faites les vérifications en mode Aperçu avant de publier. L'onglet 4 "
                  "contient le code à coller ; l'onglet 5 indique ce que devient chaque objet."),
    },
}
TOPICS = {
    "A02": ("Inventory", "Inventaire"), "A03": ("Broken references", "Références cassées"),
    "A04": ("Unused and paused objects", "Objets inutilisés ou en pause"), "A05": ("Duplicates", "Doublons"),
    "A06": ("Tag settings", "Réglages des balises"), "A07": ("Trigger conditions", "Conditions des déclencheurs"),
    "A08": ("Firing options and sequencing", "Options de déclenchement et ordre"),
    "A09": ("Consent platform", "Plateforme de consentement"), "A10": ("Consent blocking", "Blocage selon le consentement"),
    "A11": ("Consent Mode", "Consent Mode"), "A12": ("Sending to a server container", "Envoi vers un conteneur serveur"),
    "A13": ("Server container handoff", "Passage au conteneur serveur"), "A14": ("Variables", "Variables"),
    "A15": ("Google tag settings", "Réglages de la balise Google"),
    "A16": ("Google tags and page views", "Balises Google et pages vues"),
    "A17": ("GA4 events and parameters", "Événements et paramètres GA4"), "A18": ("Ecommerce", "E-commerce"),
    "A19": ("Ads and other vendors", "Publicité et autres outils"), "A20": ("Values and formats", "Valeurs et formats"),
    "A21": ("Personal data and secrets", "Données personnelles et secrets"),
    "A22": ("Custom code and templates", "Code et modèles personnalisés"),
    "A23": ("Environments", "Environnements"), "A24": ("Naming and folders", "Nommage et dossiers"),
    "A25": ("Simplification", "Simplification"), "A26": ("Target architecture", "Architecture cible"),
}
MAX_CELL = 32000
LINK_PREFIX = '=HYPERLINK("#'


def language(context: dict) -> str:
    return "fr" if str(context.get("language", "en")).lower().startswith("fr") else "en"


def is_link(value) -> bool:
    return isinstance(value, str) and value.startswith(LINK_PREFIX)


def _link(sheet: str, rows: list[int], text: str = "") -> str:
    """Sort-safe link to the first row number; the text names every row."""
    if not rows:
        return ""
    text = (text or ", ".join(f"#{r}" for r in rows)).replace('"', "'")
    return f'=HYPERLINK("#\'{sheet}\'!A"&MATCH({rows[0]},\'{sheet}\'!A:A,0),"{text}")'


def template_code(raw: dict) -> str:
    return template_sections(strip_template(raw)).get(TEMPLATE_CODE, "")


def _parts(text: str) -> list[str]:
    return [text[i:i + MAX_CELL] for i in range(0, len(text), MAX_CELL)]


class Layout:
    """Every row of the five tabs, computed once from final.json; `build` writes it and `verify` reads it back."""

    def __init__(self, model: Model, final: dict, context: dict):
        self.lang = language(context)
        self.lab = LABELS[self.lang]
        self.sheets = self.lab["sheets"]
        self.ui = Interface(model, final["operations"], self.lang)
        self.model, self.final = model, final
        self.findings = final["findings"]
        self.row_of = {f["id"]: n for n, f in enumerate(self.findings, 1)}
        self.ops = {o["id"]: o for o in final["operations"]}
        # Each operation is written once, under the first (highest-priority) finding it serves.
        self.host = {oid: min((self.row_of[f] for f in o["findings"] if f in self.row_of), default=0)
                     for oid, o in self.ops.items()}
        self.question_no: dict[str, int] = {}
        self.question_rows: list[tuple[str, list[int]]] = []
        for f in self.findings:
            text = self.ui.cite(f.get("question", ""))
            if text:
                if text not in self.question_no:
                    self.question_no[text] = len(self.question_rows) + 1
                    self.question_rows.append((text, []))
                self.question_rows[self.question_no[text] - 1][1].append(self.row_of[f["id"]])
        # Execution semantics come from the engine: each action in the state it applies to, and the order in
        # which declared prerequisites run.
        self.trace = trace(model, final["operations"])
        self.order = {o["id"]: i for i, o in enumerate(execution_order(list(self.ops.values())))}
        self.code_phrase: dict[tuple[str, int], str] = {}  # (operation, action number) -> where its code is
        self.code = self._code_rows()

    # ---------- tab 2 ----------

    def steps(self, op: dict) -> list[str]:
        return self.ui.steps(self.trace[op["id"]], lambda step: self.code_phrase.get((step.op, step.index), ""))

    def placed(self, n: int, f: dict) -> list[dict]:
        """The changes written on finding row n, prerequisites first."""
        return sorted((self.ops[o] for o in f.get("operations", []) if o in self.ops and self.host[o] == n),
                      key=lambda op: self.order[op["id"]])

    def finding_row(self, n: int, f: dict) -> list:
        lab, ui = self.lab, self.ui
        placed = self.placed(n, f)
        number = {op["id"]: i for i, op in enumerate(placed, 1)}
        elsewhere = [self.ops[o] for o in f.get("operations", []) if o in self.ops and self.host[o] != n]
        todo = [ui.cite(f["recommendation"])] if f.get("recommendation") else []
        checks, after = [], set()
        for i, op in enumerate(placed, 1):
            first = sorted(number[d] for d in op.get("depends_on", []) if d in number)
            head = lab["change"].format(n=i, title=ui.cite(op["title"]))
            block = [head + (f" ({lab['after_change'].format(n=', '.join(map(str, first)))})" if first else "")]
            if op.get("preconditions"):
                block.append(lab["before"].format(text=ui.cite(op["preconditions"])))
            block += [f"{k}. {s}" for k, s in enumerate(self.steps(op), 1)]
            todo.append("\n".join(block))
            prefix = f"{i}. " if len(placed) > 1 else ""
            checks.append(prefix + lab["check"].format(text=ui.cite(op["verification"])) + "\n" +
                          lab["undo"].format(text=ui.cite(op["rollback"])))
            after |= {self.host[d] for d in op.get("depends_on", []) if self.host.get(d) and self.host[d] != n}
        for op in elsewhere:
            todo.append(lab["done_by"].format(title=ui.cite(op["title"]), row=f"#{self.host[op['id']]}"))
        kind = f["class"]
        state = f.get("readiness", "n/a") if kind in ("defect", "optimisation") else kind
        text = ui.cite(f.get("question", ""))
        status = lab["status"][state].format(q=f"Q{self.question_no[text]}" if text in self.question_no else "")
        if after:
            status += "; " + lab["status"]["after"].format(rows=", ".join(f"#{r}" for r in sorted(after)))
        return [n, lab["priority"].get(f["severity"], f["severity"]), lab["type"][kind],
                ui.pick(TOPICS[f["area"]]) if f["area"] in TOPICS else f["area"], ui.cite(f["title"]),
                "\n".join(ui.label(k) for k in f["objects"]), ui.cite(f["situation"]), ui.cite(f["impact"]),
                "\n\n".join(todo), status, "\n\n".join(checks), "", "",
                "\n".join([f["id"]] + [op["id"] for op in placed])]

    # ---------- tab 4 ----------

    def _code_rows(self) -> list[list]:
        """Current and new code for every traced action that writes code, in execution order per object, plus each
        code object no action changes. Current is the code in the state the action applies to (after its declared
        prerequisites and its own preceding actions). A changed template is a whole template file (imported). Ref
        is the action's identity (operation · action number); long text continues on numbered part rows."""
        writes: dict[str, list] = {}
        for op_id in sorted(self.trace, key=lambda o: self.order[o]):
            for step in self.trace[op_id]:
                if step.before is not None:
                    a = step.action
                    writes.setdefault(a["ref"] if a["do"] == "create" else a["object"], []).append(step)
        keys = [o.key for o in self.model.objects.values()
                if code_of(o) or (o.layer == "template" and template_code(o.raw))]
        keys += [k for k in writes if k not in keys]
        by_object: dict[str, list[int]] = {}
        for f in self.findings:
            for k in f["objects"]:
                by_object.setdefault(k, []).append(self.row_of[f["id"]])
        rows = []
        for key in keys:
            obj = self.model.get(key)
            if obj:
                kind, status = self.ui.type_name(obj.layer, obj.type), "paused" if obj.paused else "active"
            else:
                layer, created = self.ui.created[key]
                kind, status = self.ui.type_name(layer, str(created.get("type", ""))), "proposed"
            entries = []
            for step in writes.get(key, []):
                label = f"{self.ui.at(step.names).label(key)} · {self.ui.cite(self.ops[step.op]['title'])}"
                imported = bool(obj) and obj.layer == "template"
                ref = f"{step.op} · {step.index}"
                self.code_phrase[(step.op, step.index)] = self.lab["import" if imported else "new_code"].format(
                    ref=f"{self.sheets[3]} › {label} ({self.lab['code'][7]} {ref})")
                entries.append((label, step.before, step.after, [self.host[step.op]], ref))
            if not entries:
                entries = [(self.ui.label(key), code_of(obj) or template_code(obj.raw), "", [], "")]
            for label, current, new, changed, ref in entries:
                current, new = _parts(redact(current or "")), _parts(redact(new or ""))
                total = max(len(current), len(new), 1)
                for i in range(total):
                    part = f" ({self.lab['part'].format(n=i + 1, total=total)})" if total > 1 else ""
                    rows.append([label + part, kind, self.lab["object_status"][status],
                                 _link(self.sheets[1], sorted(set(by_object.get(key, [])))),
                                 current[i] if i < len(current) else "", new[i] if i < len(new) else "",
                                 _link(self.sheets[1], changed), ref])
        return rows

    # ---------- tabs 3 and 5 ----------

    def question_table(self) -> list[list]:
        titles = {self.row_of[f["id"]]: self.ui.cite(f["title"]) for f in self.findings}
        return [[f"Q{i}", text, _link(self.sheets[1], rows, "\n".join(f"#{r} {titles[r]}" for r in rows)), "", ""]
                for i, (text, rows) in enumerate(self.question_rows, 1)]

    def inventory(self) -> list[list]:
        rows = []
        for c in self.final["coverage"]:
            folder = c.get("folder", "")
            folder = self.ui.name(folder if folder.startswith("new:") else f"folder:{folder}") if folder else ""
            linked = {self.row_of[f] for f in c["findings"] if f in self.row_of} | \
                     {self.host[o] for o in c["operations"] if self.host.get(o)}
            rows.append([self.ui.named(c["layer"], c["name"]), self.ui.type_name(c["layer"], c["type"]), folder,
                         self.lab["object_status"][c["status"]], self.lab["outcome"][c["outcome"]],
                         c.get("new_name") or "", _link(self.sheets[1], sorted(linked))])
        return rows

    # ---------- tab 1 ----------

    def summary(self) -> list[list]:
        container = self.model.cv.get("container") or {}
        s = self.final.get("summary") or {}
        h = self.lab["summary"]
        limits = self.lab["boundary"] + ("\n\n" + self.ui.cite(s["limits"]) if s.get("limits") else "")
        rows = [[h[0], f"{container.get('publicId', '')} {container.get('name', '')}".strip()],
                [h[1], str(self.model.cv.get("containerVersionId", ""))],
                [h[2], dt.date.today().isoformat() + f" · skill v{VERSION}"], [h[3], limits],
                [h[4], self.ui.cite(s.get("target", ""))], [h[5], self.ui.cite(s.get("naming", ""))],
                [h[6], str(len(self.question_rows))], [h[7], self.lab["howto"]], []]
        classes = list(self.lab["type"])
        rows.append([self.lab["grid"]] + [self.lab["type"][c] for c in classes] + [self.lab["total"]])
        for sev in ("Critical", "High", "Medium", "Low", "None"):
            counts = [sum(1 for f in self.findings if f["severity"] == sev and f["class"] == c) for c in classes]
            if sum(counts):
                rows.append([self.lab["priority"][sev]] + counts + [sum(counts)])
        return rows

    def content(self) -> dict[str, list[list]]:
        """Every cell of the workbook, by tab: what `build` writes and `verify` compares the saved file with."""
        lab, s = self.lab, self.sheets
        return {
            s[0]: self.summary(),
            s[1]: [list(lab["findings"])] + [self.finding_row(n, f) for n, f in enumerate(self.findings, 1)],
            s[2]: [list(lab["questions"])] + self.question_table(),
            s[3]: [list(lab["code"])] + self.code,
            s[4]: [list(lab["inventory"])] + self.inventory(),
        }


WIDTHS = {1: (5, 9, 13, 16, 32, 30, 48, 40, 70, 20, 42, 12, 24, 12), 2: (6, 70, 50, 50, 22),
          3: (36, 20, 10, 14, 80, 80, 12, 12), 4: (48, 26, 22, 10, 22, 36, 14)}


def build(model: Model, final: dict, context: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    lay = Layout(model, final, context)
    content = lay.content()
    wb = Workbook()
    wb.remove(wb.active)
    bold, white = Font(bold=True), Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, (title, rows) in enumerate(content.items()):
        ws = wb.create_sheet(title)
        for row in rows:
            for cell in row:
                if isinstance(cell, str) and len(cell) > MAX_CELL:
                    raise ValueError(f"{title}: a cell exceeds {MAX_CELL} characters")
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.alignment = wrap
                if isinstance(cell.value, str) and not is_link(cell.value):
                    cell.data_type = "s"
        if index == 0:
            ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 110
            for row in ws.iter_rows():
                for cell in row[:1] if row[0].value != lay.lab["grid"] else row:
                    cell.font = bold
            continue
        for cell in ws[1]:
            cell.font, cell.fill = white, fill
        for i, w in enumerate(WIDTHS[index], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "F2" if index == 1 else "B2"
        ws.auto_filter.ref = ws.dimensions
        if index == 1 and len(rows) > 1:
            choice = DataValidation(type="list", formula1=f'"{lay.lab["decisions"]}"', allow_blank=True)
            ws.add_data_validation(choice)
            choice.add(f"L2:L{len(rows)}")
            for row in ws.iter_rows(min_row=2, min_col=14, max_col=14):
                row[0].font = Font(color="808080", size=8)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def verify(path: Path, model: Model, final: dict, context: dict) -> list[str]:
    """Reopen the saved workbook: every cell equals the layout computed from final.json; independently of that
    layout, every code row holds its action's before/after code from the engine trace and no change is listed
    before its prerequisite; text stays literal (only navigation links are formulas); no email-like value is left
    unmasked."""
    from openpyxl import load_workbook

    lay = Layout(model, final, context)
    expected = lay.content()
    wb = load_workbook(path)
    problems = []
    if wb.sheetnames != list(expected):
        return [f"tabs are {wb.sheetnames}, expected {list(expected)}"]
    for title, rows in expected.items():
        ws = wb[title]
        saved = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wanted = [[str(v) for v in row] for row in rows]
        for r in range(max(len(saved), len(wanted))):
            have = saved[r] if r < len(saved) else []
            want = wanted[r] if r < len(wanted) else []
            width = max(len(have), len(want))
            if have + [""] * (width - len(have)) != want + [""] * (width - len(want)):
                problems.append(f"{title} row {r + 1} differs from final.json")
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if cell.data_type == "f" and not is_link(v):
                    problems.append(f"{title}!{cell.coordinate} is stored as a formula, not text")
                if isinstance(v, str) and any(not m.lower().endswith((".png", ".jpg", ".gif", ".svg"))
                                              for m in EMAIL_RE.findall(v)):
                    problems.append(f"{title}!{cell.coordinate}: unmasked email-like value")
    # Independent of the layout: saved code rows, read by their Ref, against a fresh engine trace.
    saved: dict[str, list[str]] = {}
    for row in wb[lay.sheets[3]].iter_rows(min_row=2, max_col=8, values_only=True):
        row = tuple(row) + (None,) * (8 - len(row))
        if row[7]:
            current, new = saved.setdefault(str(row[7]), ["", ""])
            saved[str(row[7])] = [current + str(row[4] or ""), new + str(row[5] or "")]
    expected = {f"{s.op} · {s.index}": [redact(s.before or ""), redact(s.after or "")]
                for steps in trace(model, final["operations"]).values() for s in steps if s.before is not None}
    for ref in sorted(set(expected) | set(saved)):
        if saved.get(ref) != expected.get(ref):
            problems.append(f"{lay.sheets[3]}: the saved code for {ref} is not that action's own before/after code")
    # Independent of the layout: within each finding row, no change is listed before one of its prerequisites.
    ops = {o["id"]: o for o in final["operations"]}

    def needs(op_id: str) -> set[str]:
        out, pending = set(), list(ops[op_id].get("depends_on", []))
        while pending:
            dep = pending.pop()
            if dep in ops and dep not in out:
                out.add(dep)
                pending += ops[dep].get("depends_on", [])
        return out
    for row in wb[lay.sheets[1]].iter_rows(min_row=2, max_col=14, values_only=True):
        row = tuple(row) + (None,) * (14 - len(row))
        listed = [o for o in str(row[13] or "").split("\n")[1:] if o in ops]
        for i, op_id in enumerate(listed):
            if needs(op_id) & set(listed[i + 1:]):
                problems.append(f"{lay.sheets[1]} row {row[0]}: {op_id} is listed before its prerequisite")
    return problems
